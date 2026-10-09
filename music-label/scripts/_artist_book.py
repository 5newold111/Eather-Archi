"""
アーティスト管理表（スプレッドシート）の読み書き部品。

管理表はアーティストの情報の「正本」。オーナーが書き足し、機械は読むだけ（提案は「状態＝提案」として足す）。
  ・Google スプレッドシート … .env に ARTIST_BOOK_GSHEET_ID があればこちらを読む（正本）
        GOOGLE_SERVICE_ACCOUNT_FILE（サービスアカウントの鍵）もあれば書き込みもできる。
        鍵が無ければ「リンクを知っている全員が閲覧可」の共有で読み込みだけ行い、機械の提案は
        out/artist_book/pending.xlsx に貯める（Google の表へは人が貼るか、鍵を設定する）
  ・Excel ファイル（music-label/artists.xlsx）… Google の設定が無いときの正本。デザインの見本も兼ねる

タブには 2 つの向きがある
  縦型（1 列 = 1 組・1 人）：アーティスト／メンバー／声と曲づくり／見た目 … 項目が多いので、組を横に並べて見比べる
  横型（1 行 = 1 件）　　　：人生年表・デビューの経緯・傷とタトゥー・影響・歌の種・歌詞の試作・近況・参考曲の候補・変更の提案
読むときは向きを自動で見分けるので、古い横型の表もそのまま読める。
"""
from __future__ import annotations

import io
import json
import os
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path

from _common import OUT, ROOT, ssl_context

STATUS = ["提案", "採用", "保留", "却下", "反映済み"]
ADOPTED = {"採用", ""}

# (キー, 見出し, 種類) 種類：text / int / list（「、」か改行で区切る）/ bool（○ で真）
TABS: dict[str, tuple[str, list[tuple[str, str, str]]]] = {
    "acts": ("アーティスト", [
        ("name", "名前", "text"), ("status", "状態", "text"), ("id", "ID", "text"),
        ("reading", "読み方", "text"), ("name_meaning", "名前の由来", "text"), ("formation", "構成", "text"),
        ("label", "レーベル", "text"), ("updated", "更新日", "text"),
        ("origin", "出身の設定", "text"), ("base_now", "今の拠点", "text"), ("culture", "文化", "text"),
        ("scene", "聴かれる場面", "text"), ("strengths", "強み", "list"),
        ("debut_summary", "デビューの経緯（まとめ）", "text"), ("fan_growth", "ファンのつき方", "text"),
        ("expression", "音楽以外の表現", "text"),
        ("dynamics", "人間関係・力関係", "text"), ("lyricist", "作詞の担当", "text"), ("composer", "作曲の担当", "text"),
        ("talk", "普段の話題", "text"),
        ("values", "人生観・価値観（組として）", "text"), ("future", "今後どうしていきたいか", "text"),
        ("current_chapter", "今の章（いまの暮らしと気持ち）", "text"), ("landscape_words", "歌によく出る風景・言葉", "list"),
        ("bio_en", "英語の紹介文", "text"), ("note", "メモ", "text")]),
    "members": ("メンバー", [
        ("name", "名前", "text"), ("status", "状態", "text"), ("act_id", "アーティストID", "text"), ("id", "メンバーID", "text"),
        ("role", "担当", "text"), ("age", "年齢", "int"), ("birthday", "誕生日", "text"),
        ("gender", "性別", "text"), ("birthplace", "出身地", "text"), ("roots", "ルーツ", "text"),
        ("languages", "話せる言葉", "list"),
        ("height", "身長(cm)", "int"), ("weight", "体重(kg)", "int"),
        ("hair_color", "髪の色", "text"), ("hair_length", "髪の長さ", "text"), ("hair_style", "髪型", "text"),
        ("fashion", "好きなファッション", "text"),
        ("personality", "性格", "text"), ("values", "人生観・価値観", "text"),
        ("speech", "話し方・訛り", "text"), ("catchphrases", "口癖", "list"), ("endings", "語尾", "text"),
        ("family", "家族構成", "text"), ("friends", "友人", "text"), ("holidays", "休日の過ごし方", "text"),
        ("expression", "音楽以外の表現", "text"),
        ("likes", "好きなこと", "list"), ("like_words", "好きな言葉", "list"),
        ("like_colors", "好きな色", "list"), ("like_foods", "好きな食べ物", "list"),
        ("dislikes", "嫌いなこと", "list"), ("dislike_words", "嫌いな言葉", "list"),
        ("dislike_colors", "嫌いな色", "list"), ("dislike_foods", "嫌いな食べ物", "list"), ("note", "メモ", "text")]),
    "timeline": ("人生年表", [
        ("act_id", "アーティストID", "text"), ("member_id", "メンバーID", "text"), ("status", "状態", "text"),
        ("age_from", "年齢（から）", "int"), ("age_to", "年齢（まで）", "int"), ("years", "西暦", "text"),
        ("place", "場所", "text"), ("life", "暮らし・学校", "text"), ("people", "周りにいた人", "text"),
        ("event", "出来事", "text"), ("feeling", "そのときの気持ち", "text"), ("seed", "歌の種にする", "bool")]),
    "debut": ("デビューの経緯", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"), ("when", "時期", "text"),
        ("venue", "場（路上・ライブハウス・配信・大会・学園祭・SNS など）", "text"), ("what", "やったこと", "text"),
        ("reaction", "反応・ファンの数", "text"), ("turning_point", "転機", "text")]),
    "marks": ("傷・タトゥー", [
        ("act_id", "アーティストID", "text"), ("member_id", "メンバーID", "text"), ("status", "状態", "text"),
        ("kind", "種類", "text"), ("where", "体の場所", "text"), ("when", "いつ", "text"),
        ("what", "図柄・内容", "text"), ("meaning", "込めた想い", "text")]),
    "influences": ("影響", [
        ("act_id", "アーティストID", "text"), ("member_id", "メンバーID", "text"), ("status", "状態", "text"),
        ("kind", "種類（音楽・本・絵・映画・言葉）", "text"), ("title", "名前・題名", "text"),
        ("author", "作者・アーティスト", "text"), ("taken", "受け取ったもの", "text")]),
    "seeds": ("歌の種", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"), ("id", "種ID", "text"),
        ("period", "時期", "text"), ("feeling", "気持ち", "text"), ("scene", "情景", "text"), ("pov", "視点", "text"),
        ("idea", "歌のアイデア", "text"), ("words", "言葉の候補", "list")]),
    "lyrics": ("歌詞の試作", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"), ("seed_id", "種ID", "text"),
        ("title", "仮タイトル", "text"), ("chorus", "サビ（英語）", "text"), ("verse", "1番の書き出し（英語）", "text"),
        ("meaning_ja", "日本語の意味と気持ち", "text")]),
    "updates": ("近況", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"), ("date", "日付", "text"),
        ("event", "出来事", "text"), ("feeling", "気持ち", "text"), ("to_song", "歌にする", "bool")]),
    "choices": ("参考曲の候補", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"), ("title", "曲名", "text"),
        ("artist", "アーティスト", "text"), ("year", "年", "text"), ("why", "この組が選ぶ理由", "text"),
        ("slots", "借りたい要素", "list"), ("acquired", "入手", "bool"), ("analyzed", "解析", "bool")]),
    "music": ("声と曲づくり", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"),
        ("v_sex", "声の性別（女性・男性・男女・なし）", "text"), ("v_range", "声域（高・中・低）", "text"),
        ("fixed_timbre", "声質を固定", "bool"), ("low", "楽に出る一番下", "text"), ("high", "楽に出る一番上", "text"),
        ("chest_top", "地声の上限", "text"), ("switch", "裏声に切り替わる音", "text"), ("falsetto", "裏声の質", "text"),
        ("v_strength", "声の強み", "text"), ("v_weakness", "声の弱み", "text"),
        ("primary_tech", "一番得意な歌い方", "text"), ("secondary_tech", "ほかの得意技", "list"),
        ("avoid_tech", "避ける歌い方", "list"), ("wordless", "言葉のないパート", "text"),
        ("tic1", "毎曲の癖1", "text"), ("tic1_where", "癖1の場所", "text"),
        ("tic2", "毎曲の癖2", "text"), ("tic2_where", "癖2の場所", "text"),
        ("tic3", "毎曲の癖3", "text"), ("tic3_where", "癖3の場所", "text"),
        ("guest", "客演するときの声の説明", "text"),
        ("melodic", "旋律の動き", "text"), ("chorus_peak", "サビの頂点", "text"), ("verse_register", "A メロの音域", "text"),
        ("phrase_endings", "フレーズの終わり方", "text"),
        ("genres", "ジャンル", "list"), ("palette", "楽器の色", "list"), ("bpm_min", "BPM 最低", "int"),
        ("bpm_max", "BPM 最高", "int"), ("dna", "DNA タグ", "list"),
        ("modulation", "転調", "text"), ("intro", "イントロ", "text"), ("outro", "アウトロ", "text"), ("structure", "構成", "text"),
        ("inst_rule", "楽器のルール", "text"), ("rhythm", "リズム", "text"), ("always", "毎曲必ず", "list"),
        ("never", "使わない", "list"),
        ("themes", "歌詞のテーマ", "list"), ("pov", "歌詞の視点", "text"),
        ("lang", "歌詞の言語", "text"), ("tone", "歌詞のトーン", "text"), ("avoid_topics", "避ける話題", "list"),
        ("lyrics_tic", "歌詞の癖", "text"), ("lyric_mode", "歌詞の作り方（核を固定・全部書く・お題だけ・歌なし）", "text"),
        ("partners", "相性のよい組", "list"), ("collab_style", "コラボの形", "text")]),
    "visual": ("見た目", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"), ("space", "空間", "text"), ("light", "光", "text"),
        ("materials", "素材", "list"), ("palette", "色", "list"), ("camera", "カメラ", "text"),
        ("cover_rule", "ジャケットの決まり", "text"), ("conceal", "顔を見せない方法", "text"),
        ("conceal2", "副の方法", "list"), ("never", "絶対にしないこと", "list"), ("logo_concept", "ロゴの考え方", "text"),
        ("logo_mark", "ロゴの形", "text"), ("logo_color", "ロゴの色", "text"), ("logo_font", "ロゴの書体", "text"),
        ("taste", "写真の質感", "text"), ("composition", "構図", "text"), ("location", "撮影場所", "text"),
        ("brightness", "明るさ", "text"), ("contrast", "コントラスト", "text"), ("temperature", "色温度", "text"),
        ("background", "背景", "text"), ("position", "立ち位置", "text"), ("gesture", "しぐさ", "text")]),
    "changes": ("変更の提案", [
        ("act_id", "アーティストID", "text"), ("status", "状態", "text"), ("date", "日付", "text"),
        ("tab", "タブ", "text"), ("field", "項目", "text"), ("current", "今の値", "text"),
        ("proposed", "提案する値", "text"), ("reason", "理由", "text")]),
}
TAB_ORDER = list(TABS)
VERTICAL = {"acts", "members", "music", "visual"}          # 1 列 = 1 組（1 人）
KEY_FIELD = {"acts": "id", "members": "id"}                  # 縦型で「どの組・どの人か」を決める行。無ければ act_id
SPLIT = re.compile(r"[、\n]+")

# 縦型タブの見出し帯（■ の行）。ここに無い項目は「その他」に入る
SECTIONS: dict[str, list[tuple[str, list[str]]]] = {
    "acts": [("", ["name", "status", "id"]),
             ("基本", ["reading", "name_meaning", "formation", "label", "updated"]),
             ("生まれと暮らし", ["origin", "base_now", "culture"]),
             ("活動", ["scene", "strengths", "debut_summary", "fan_growth", "expression"]),
             ("人と関係", ["dynamics", "lyricist", "composer", "talk"]),
             ("考え方といま", ["values", "future", "current_chapter", "landscape_words"]),
             ("公開用", ["bio_en"]), ("メモ", ["note"])],
    "members": [("", ["name", "status", "act_id", "id"]),
                ("基本", ["role", "age", "birthday", "gender", "birthplace", "roots", "languages"]),
                ("見た目", ["height", "weight", "hair_color", "hair_length", "hair_style", "fashion"]),
                ("内面", ["personality", "values"]),
                ("話し方", ["speech", "catchphrases", "endings"]),
                ("暮らしと人", ["family", "friends", "holidays", "expression"]),
                ("好きなもの", ["likes", "like_words", "like_colors", "like_foods"]),
                ("嫌いなもの", ["dislikes", "dislike_words", "dislike_colors", "dislike_foods"]),
                ("メモ", ["note"])],
    "music": [("", ["act_id", "status"]),
              ("声", ["v_sex", "v_range", "fixed_timbre", "low", "high", "chest_top", "switch", "falsetto",
                     "v_strength", "v_weakness"]),
              ("歌い方", ["primary_tech", "secondary_tech", "avoid_tech", "wordless", "guest"]),
              ("毎曲の癖", ["tic1", "tic1_where", "tic2", "tic2_where", "tic3", "tic3_where"]),
              ("旋律", ["melodic", "chorus_peak", "verse_register", "phrase_endings"]),
              ("音", ["genres", "palette", "bpm_min", "bpm_max", "dna"]),
              ("曲の形", ["modulation", "intro", "outro", "structure", "inst_rule", "rhythm", "always", "never"]),
              ("歌詞", ["themes", "pov", "lang", "tone", "avoid_topics", "lyrics_tic", "lyric_mode"]),
              ("コラボ", ["partners", "collab_style"])],
    "visual": [("", ["act_id", "status"]),
               ("空間", ["space", "light", "materials", "palette", "camera"]),
               ("ジャケット", ["cover_rule"]),
               ("顔を見せない", ["conceal", "conceal2", "never"]),
               ("ロゴ", ["logo_concept", "logo_mark", "logo_color", "logo_font"]),
               ("写真", ["taste", "composition", "location", "brightness", "contrast", "temperature", "background",
                        "position", "gesture"])],
}

# 選ぶ欄（プルダウン）
CHOICES: dict[tuple[str, str], list[str]] = {
    ("acts", "formation"): ["ソロ", "デュオ", "バンド", "ボーカルグループ", "プロデューサー"],
    ("acts", "label"): ["本体", "sleep", "morning", "focus", "move"],
    ("members", "gender"): ["女性", "男性", "ノンバイナリー"],
    ("marks", "kind"): ["傷", "タトゥー", "ピアス", "なし"],
    ("influences", "kind"): ["音楽", "本", "絵", "映画", "言葉"],
    ("music", "v_sex"): ["女性", "男性", "男女", "なし"],
    ("music", "v_range"): ["高", "中", "低"],
    ("music", "lang"): ["英語", "日本語", "スペイン語", "韓国語", "ポルトガル語", "なし"],
    ("music", "lyric_mode"): ["核を固定", "全部書く", "お題だけ", "歌なし"],
}

# 見出しに付ける説明（Google スプレッドシートでは「メモ」として、セルに触れると出る）
HELP = {
    ("acts", "id"): "英小文字と _ だけ（例 nao_easterly）。ほかのタブの行とつなぐ目印。あとから変えない",
    ("acts", "scene"): "この組の曲が聴かれる場面。毎週の曲づくりの前提になる",
    ("acts", "current_chapter"): "いまの暮らしと気持ち。ここを書き換えると、次の曲の空気が変わる",
    ("acts", "bio_en"): "ストアや SNS のプロフィールに使う英語の紹介文（3〜4 文）",
    ("members", "act_id"): "どの組のメンバーか（アーティストのタブの ID）",
    ("timeline", "seed"): "○ を付けた時期は、曲のテーマの候補として作曲の指示に渡る",
    ("updates", "to_song"): "○ を付けると、まだ歌にしていない新しい出来事から順に、毎週の曲のテーマになる",
    ("choices", "slots"): "lyrics・worldview・instruments・performance・structure・harmony・groove・phrase・vocal から 1〜3 個",
    ("choices", "acquired"): "正規に購入して手元にあれば ○。解析すると、この組の曲づくりで優先される",
    ("music", "dna"): "曲の特徴を表す短い英語のタグ（例 morning, synth, uplift）。参考曲選びに使う",
    ("music", "tic1"): "毎曲必ず入れる歌い方の癖。Suno のタグで書ける具体さで（例：サビ前に裏声の ooh を 1 本）",
    ("music", "lyric_mode"): "核を固定＝サビなどをこちらで決め、残りを Suno に／全部書く／お題だけ＝歌詞を Suno に任せる／歌なし",
    ("music", "partners"): "相性のよい組の ID（例 hollowplan）",
    ("music", "collab_style"): "リミックス役の組は、ここに remix の語を入れる（5 週に 1 回のリミックスの担当になる）",
    ("visual", "conceal"): "顔は絶対に見せない。逆光・後ろ姿・手元など",
}

# ---------------------------------------------------------------------------
# デザイン
# ---------------------------------------------------------------------------
FONT = "Arial"
INK, SUB, LINE = "2B2B2B", "8A8A8A", "D9D6CF"
GROUPS = {   # タブの色分け：(見出しの色, 帯の色, 項目欄の色, タブの色)
    "person": ("EDE5D6", "9C7F57", "F8F4EC", "B89B72"),
    "song":   ("DDEBE1", "5E8B6E", "F2F7F3", "6F9A7E"),
    "look":   ("DFE7EE", "62809A", "F2F5F8", "7D97AE"),
    "ops":    ("E7E7E7", "7A7A7A", "F5F5F5", "9E9E9E"),
}
TAB_GROUP = {"acts": "person", "members": "person", "timeline": "person", "debut": "person", "marks": "person",
             "influences": "person", "seeds": "song", "lyrics": "song", "updates": "song", "choices": "song",
             "music": "song", "visual": "look", "changes": "ops"}
STATUS_STYLE = {"提案": ("FFF1C2", "7A5A00"), "採用": ("D9EFDE", "1E6B3A"), "保留": ("E8ECEF", "52626E"),
                "却下": ("F7D9D9", "A12B2B"), "反映済み": ("DCE6F6", "1F4E99")}
ACT_COLORS = ["F3DDB0", "CAD5EE", "CFE6D4", "D6E0E6", "E4D7EF", "F5D2C6", "DCE9C9", "F0E4BC", "CDE4EC", "EAD4DF"]
WIDE = 46        # 縦型の 1 組ぶんの列幅
NARROW = {"act_id": 17, "member_id": 12, "status": 9, "id": 12, "seed_id": 11, "age_from": 10, "age_to": 10,
          "years": 13, "date": 12, "when": 18, "year": 7, "kind": 12, "where": 16, "place": 26, "period": 22,
          "venue": 26, "title": 26, "author": 22, "tab": 14, "field": 20, "seed": 10, "to_song": 9,
          "acquired": 7, "analyzed": 7, "pov": 18}

INTRO_TITLE = "EtherArchi アーティスト管理表"
INTRO_LEAD = ("所属アーティストの人生と、曲づくりの設定をまとめた正本です。書き足すと、次の曲づくりから使われます"
              "（毎週火曜の朝に機械が読み込みます）。")
INTRO_RULES = [
    "状態は、提案＝機械の案／採用＝使う／保留＝あとで決める／却下＝使わない。空欄は採用として扱います。機械は『提案』を足すだけで、採用はしません。",
    "アーティスト・メンバー・声と曲づくり・見た目のタブは縦型です。1 列が 1 組（1 人）で、組を横に並べて見比べられます。新しい組は右の空いた列に書きます。",
    "人生年表・デビューの経緯・傷とタトゥー・影響・歌の種・歌詞の試作・近況・参考曲の候補は横型です。1 つの出来事・1 項目を 1 行で、何行でも足せます。",
    "組とのつながりは ID（英小文字と _。例 nao_easterly）で、メンバーとのつながりはメンバーID で書きます。",
    "複数の値は「、」か改行で区切ります（例：好きな色 → 朝の水色、生成り）。○ は『はい』の意味です。",
    "実在のアーティスト名・曲名・本の題名は『影響』と『参考曲の候補』のタブにだけ書きます。ほかの欄に書くと、作曲の指示文に混ざるおそれがあります。",
    "毎週の曲のテーマは『近況』で『歌にする』に ○ が付いた新しい出来事 → 『歌の種』の順に、まだ使っていないものから選ばれます。",
    "見出し（項目名）の文字は変えないでください。機械は見出しの文字で読みます。列や行の順番を入れ替えるのは大丈夫です。",
    "灰色の小さな文字の ID は機械用の目印です。見出しにカーソルを置くと、書き方の説明が出るものがあります。",
]
TAB_HELP = {
    "acts": "1 列 = 1 組。名前・出身・デビューの経緯・人間関係・価値観・今の章",
    "members": "1 列 = 1 人。年齢・見た目・性格・話し方・家族・好き嫌い",
    "timeline": "生まれてから今と、これから。1 行 = 1 つの時期",
    "debut": "配信の前に、どこでどう人が集まったか",
    "marks": "傷とタトゥー（無い人は『なし』と理由）",
    "influences": "影響を受けた実在の音楽・本・絵・映画・言葉",
    "seeds": "人生の出来事から生まれた歌のアイデア",
    "lyrics": "歌詞の試作（英語）と日本語の意味",
    "updates": "最近の出来事と気持ち。『歌にする』○ で次の曲のテーマに",
    "choices": "この組が選びそうな実在の参考曲。入手・解析の印",
    "music": "1 列 = 1 組。声・歌い方・癖・音・曲の形・歌詞の決まり",
    "visual": "1 列 = 1 組。部屋・光・色・顔を見せない方法・ロゴ・写真",
    "changes": "方針転換などで機械が出す変更の案",
}


def header_to_key(tab: str) -> dict[str, tuple[str, str]]:
    return {h: (k, kind) for k, h, kind in TABS[tab][1]}


def key_to_header(tab: str) -> dict[str, str]:
    return {k: h for k, h, _ in TABS[tab][1]}


def kind_of(tab: str, key: str) -> str:
    return next(kind for k, _, kind in TABS[tab][1] if k == key)


def to_cell(value, kind: str):
    if value is None:
        return ""
    if kind == "list":
        return "、".join(str(v) for v in value) if isinstance(value, (list, tuple)) else str(value)
    if kind == "bool":
        return "○" if value else ""
    return value


def from_cell(value, kind: str):
    if value is None:
        value = ""
    if kind == "list":
        return [x.strip() for x in SPLIT.split(str(value)) if x.strip()]
    if kind == "bool":
        return str(value).strip() in ("○", "◯", "TRUE", "True", "true", "1", "yes", "はい")
    if kind == "int":
        try:
            return int(float(str(value).replace("cm", "").replace("kg", "").strip()))
        except ValueError:
            return None
    return str(value).strip() if not isinstance(value, (int, float)) else value


def section_layout(tab: str) -> list[tuple[str, list[str]]]:
    """縦型タブの並び（見出し帯, 項目）。SECTIONS に無い項目は最後に『その他』として足す"""
    secs = [(t, list(ks)) for t, ks in SECTIONS[tab]]
    placed = {k for _, ks in secs for k in ks}
    rest = [k for k, _, _ in TABS[tab][1] if k not in placed]
    if rest:
        secs.append(("その他", rest))
    return secs


def col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


# ---------------------------------------------------------------------------
# 読み取り（向きは自動判定）
# ---------------------------------------------------------------------------
def parse_grid(tab: str, grid: list[list]) -> tuple[list[dict], dict]:
    """セルの格子（行のリスト）→ 行の辞書のリストと、書き込み位置の情報"""
    m = header_to_key(tab)
    grid = [list(r) for r in grid]

    def s(v):
        return str(v).strip() if v is not None else ""
    # 横型：上 5 行のどこかに見出しが 3 つ以上並ぶ行がある
    for hi, row in enumerate(grid[:5]):
        if sum(1 for v in row if s(v) in m) >= 3:
            heads = [s(v) for v in row]
            items = []
            for r in range(hi + 1, len(grid)):
                vals = grid[r]
                if not any(s(v) for v in vals):
                    continue
                d = {"_pos": r + 1}
                for c, h in enumerate(heads):
                    if h in m:
                        k, kind = m[h]
                        d[k] = from_cell(vals[c] if c < len(vals) else "", kind)
                items.append(d)
            return items, {"orient": "row", "header_row": hi + 1, "heads": heads}
    # 縦型：A 列に見出しが縦に並ぶ
    rows = {s(r[0]): i for i, r in enumerate(grid) if r and s(r[0]) in m}
    if not rows:
        return [], {"orient": "col" if tab in VERTICAL else "row", "rows": {}, "width": 0}
    width = max(len(r) for r in grid)
    items = []
    for c in range(1, width):
        vals = {h: (grid[i][c] if c < len(grid[i]) else None) for h, i in rows.items()}
        if not any(s(v) for v in vals.values()):
            continue
        d = {"_pos": c + 1}
        for h, v in vals.items():
            k, kind = m[h]
            d[k] = from_cell(v, kind)
        items.append(d)
    return items, {"orient": "col", "rows": {m[h][0]: i + 1 for h, i in rows.items()}, "width": width}


def grid_of(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


class ReadOnlyBook(Exception):
    """Google スプレッドシートを鍵なし（閲覧だけ）で読んでいるときの書き込み"""


# ---------------------------------------------------------------------------
# Excel ファイル（作成とデザインもここ）
# ---------------------------------------------------------------------------
class XlsxBook:
    def __init__(self, path: Path) -> None:
        self.path = path

    def describe(self) -> str:
        return f"Excel ファイル {self.path.relative_to(ROOT) if self.path.is_relative_to(ROOT) else self.path}"

    def create(self, data: dict[str, list[dict]] | None = None) -> None:
        """デザイン付きで丸ごと書き出す（data があれば中身ごと）"""
        write_designed(self.path, data or {t: [] for t in TAB_ORDER})

    def read(self) -> dict[str, list[dict]]:
        from openpyxl import load_workbook
        return read_workbook(load_workbook(self.path, data_only=True))

    def append(self, tab: str, rows: list[dict]) -> None:
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        ws = wb[TABS[tab][0]]
        items, lay = parse_grid(tab, grid_of(ws))
        acts = act_colors(wb)
        if tab == "acts":   # 新しい組の色を先に決めておく
            for d in rows:
                acts.setdefault(d.get("id"), ACT_COLORS[len(acts) % len(ACT_COLORS)])
        for d in rows:
            if lay["orient"] == "row":
                r = (items[-1]["_pos"] if items else lay.get("header_row", 1)) + 1
                write_row(ws, tab, r, d, lay["heads"], acts)
                items.append({"_pos": r})
            else:
                c = max([i["_pos"] for i in items] + [1]) + 1
                write_col(ws, tab, c, d, lay["rows"], acts)
                items.append({"_pos": c})
        wb.save(self.path)

    def update(self, tab: str, pos: int, key: str, value) -> None:
        self.update_many([(tab, pos, key, value)])

    def update_many(self, edits: list[tuple[str, int, str, object]]) -> None:
        """(タブ, 位置, 項目, 値) をまとめて書き込む。位置は横型なら行番号、縦型なら列番号"""
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        lays: dict[str, dict] = {}
        for tab, pos, key, value in edits:
            ws = wb[TABS[tab][0]]
            if tab not in lays:
                lays[tab] = parse_grid(tab, grid_of(ws))[1]
            lay = lays[tab]
            v = to_cell(value, kind_of(tab, key))
            if lay["orient"] == "row":
                ws.cell(row=pos, column=lay["heads"].index(key_to_header(tab)[key]) + 1, value=v)
            else:
                ws.cell(row=lay["rows"][key], column=pos, value=v)
        wb.save(self.path)


def read_workbook(wb) -> dict[str, list[dict]]:
    out = {}
    for tab in TAB_ORDER:
        title = TABS[tab][0]
        out[tab] = parse_grid(tab, grid_of(wb[title]))[0] if title in wb.sheetnames else []
    return out


def act_colors(wb_or_data) -> dict[str, str]:
    """組ごとの色（アーティストのタブに出てくる順）"""
    if isinstance(wb_or_data, dict):
        ids = [a.get("id") for a in wb_or_data.get("acts", [])]
    else:
        title = TABS["acts"][0]
        ids = [a.get("id") for a in parse_grid("acts", grid_of(wb_or_data[title]))[0]] if title in wb_or_data.sheetnames else []
    return {i: ACT_COLORS[n % len(ACT_COLORS)] for n, i in enumerate(x for x in ids if x)}


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color=LINE)
    return {
        "body": Font(name=FONT, size=10, color=INK),
        "muted": Font(name=FONT, size=9, color=SUB),
        "bold": Font(name=FONT, size=10, bold=True, color=INK),
        "name": Font(name=FONT, size=13, bold=True, color=INK),
        "sec": Font(name=FONT, size=10, bold=True, color="FFFFFF"),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "center": Alignment(horizontal="center", vertical="top", wrap_text=True),
        "mid": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "box": Border(bottom=thin, right=thin),
        "fill": lambda c: PatternFill("solid", fgColor=c),
    }


def _value_style(cell, tab: str, key: str, st: dict) -> None:
    kind = kind_of(tab, key)
    cell.font = st["muted"] if key in ("id", "act_id", "member_id", "seed_id") else st["body"]
    cell.alignment = st["center"] if kind in ("int", "bool") or key == "status" else st["wrap"]
    cell.border = st["box"]


def write_row(ws, tab: str, r: int, d: dict, heads: list[str], acts: dict[str, str]) -> None:
    st = _styles()
    m = header_to_key(tab)
    for c, h in enumerate(heads, 1):
        if h not in m:
            continue
        k, kind = m[h]
        cell = ws.cell(row=r, column=c, value=to_cell(d.get(k), kind))
        _value_style(cell, tab, k, st)
        if k == "act_id" and d.get("act_id") in acts:
            cell.fill = st["fill"](acts[d["act_id"]])
            cell.font = st["bold"]


def write_col(ws, tab: str, c: int, d: dict, rows: dict[str, int], acts: dict[str, str]) -> None:
    st = _styles()
    ws.column_dimensions[col_letter(c)].width = WIDE
    act = d.get("id") if tab == "acts" else d.get("act_id")
    color = acts.get(act) or ACT_COLORS[(c - 2) % len(ACT_COLORS)]
    top = min(rows.values()) if rows else 1
    for k, r in rows.items():
        cell = ws.cell(row=r, column=c, value=to_cell(d.get(k), kind_of(tab, k)))
        _value_style(cell, tab, k, st)
        if r == top:   # いちばん上の行（名前・アーティストID）は組の色で
            cell.fill = st["fill"](color)
            cell.font = st["name"] if k == "name" else st["bold"]
            cell.alignment = st["mid"]
        elif k == "act_id" and act in acts:
            cell.fill = st["fill"](acts[act])


def write_designed(path: Path, data: dict[str, list[dict]]) -> None:
    """管理表を、デザイン付きで丸ごと書き出す（中身は data）"""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    st = _styles()
    acts = act_colors(data)
    wb = Workbook()
    intro = wb.active
    intro.title = "はじめに"
    counts: dict[str, str] = {}

    def status_rules(ws, rng: str) -> None:
        for s_, (bg, fg) in STATUS_STYLE.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{s_}"'], fill=st["fill"](bg),
                                                          font=Font(name=FONT, bold=True, color=fg)))

    def add_list(ws, values: list[str], rng: str) -> None:
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "一覧から選んでください（空欄も可）", "選べる値"
        ws.add_data_validation(dv)
        dv.add(rng)

    for tab in TAB_ORDER:
        title, cols = TABS[tab]
        head_c, band_c, label_c, tab_c = GROUPS[TAB_GROUP[tab]]
        ws = wb.create_sheet(title)
        ws.sheet_properties.tabColor = tab_c
        ws.sheet_view.showGridLines = False
        recs = data.get(tab, [])
        if tab in VERTICAL:
            # ---- 縦型：A 列に項目、1 列 = 1 組 ----
            ws.column_dimensions["A"].width = 26
            rows: dict[str, int] = {}
            r = 1
            span = max(len(recs), 5) + 1
            for sec, keys in section_layout(tab):
                if sec:
                    for c in range(1, span + 1):
                        ws.cell(row=r, column=c).fill = st["fill"](band_c)
                    ws.cell(row=r, column=1, value=f"■ {sec}").font = st["sec"]
                    ws.row_dimensions[r].height = 20
                    r += 1
                for k in keys:
                    cell = ws.cell(row=r, column=1, value=key_to_header(tab)[k])
                    cell.font = st["bold"] if r == 1 or k == "status" else st["body"]
                    cell.fill = st["fill"](head_c if r == 1 else label_c)
                    cell.alignment = st["mid"] if r == 1 else st["wrap"]
                    cell.border = st["box"]
                    if (tab, k) in HELP:
                        cell.comment = Comment(HELP[(tab, k)], "EtherArchi")
                    rows[k] = r
                    r += 1
            ws.row_dimensions[1].height = 28
            for n, d in enumerate(recs):
                write_col(ws, tab, n + 2, d, rows, acts)
            for n in range(len(recs), 5):   # 空いた列も同じ幅に（新しい組を書きやすく）
                ws.column_dimensions[col_letter(n + 2)].width = WIDE
            last_col = col_letter(len(recs) + 21)
            status_rules(ws, f"B{rows['status']}:{last_col}{rows['status']}")
            for (t, k), vals in CHOICES.items():
                if t == tab:
                    add_list(ws, vals, f"B{rows[k]}:{last_col}{rows[k]}")
            for k, _, kind in cols:
                if kind == "bool":
                    add_list(ws, ["○"], f"B{rows[k]}:{last_col}{rows[k]}")
            add_list(ws, STATUS, f"B{rows['status']}:{last_col}{rows['status']}")
            ws.freeze_panes = "B3"
            key = KEY_FIELD.get(tab, "act_id")
            counts[tab] = f"=COUNTA('{title}'!{rows[key]}:{rows[key]})-1"
        else:
            # ---- 横型：1 行目が見出し、1 行 = 1 件 ----
            heads = [h for _, h, _ in cols]
            keys = [k for k, _, _ in cols]
            for c, (k, h, kind) in enumerate(cols, 1):
                cell = ws.cell(row=1, column=c, value=h)
                cell.font = st["bold"]
                cell.fill = st["fill"](head_c)
                cell.alignment = st["mid"]
                cell.border = st["box"]
                ws.column_dimensions[col_letter(c)].width = NARROW.get(k, 40 if kind != "list" else 30)
                if (tab, k) in HELP:
                    cell.comment = Comment(HELP[(tab, k)], "EtherArchi")
            ws.row_dimensions[1].height = 32
            for n, d in enumerate(recs):
                write_row(ws, tab, n + 2, d, heads, acts)
            last = len(recs) + 500
            end = col_letter(len(cols))
            sc = col_letter(keys.index("status") + 1)
            status_rules(ws, f"{sc}2:{sc}{last}")
            ws.conditional_formatting.add(f"A2:{end}{last}", FormulaRule(
                formula=[f'${sc}2="却下"'], font=Font(name=FONT, color="A9A9A9", strike=True)))
            ws.conditional_formatting.add(f"B2:{end}{last}", FormulaRule(
                formula=['AND($A2<>"",MOD(ROW(),2)=0)'], fill=st["fill"]("FAF8F4")))
            for act, color in acts.items():
                ws.conditional_formatting.add(f"A2:A{last}", CellIsRule(operator="equal", formula=[f'"{act}"'],
                                                                        fill=st["fill"](color),
                                                                        font=Font(name=FONT, bold=True, color=INK)))
            for (t, k), vals in CHOICES.items():
                if t == tab:
                    c = col_letter(keys.index(k) + 1)
                    add_list(ws, vals, f"{c}2:{c}{last}")
            for i, (k, _, kind) in enumerate(cols, 1):
                if kind == "bool":
                    add_list(ws, ["○"], f"{col_letter(i)}2:{col_letter(i)}{last}")
            add_list(ws, STATUS, f"{sc}2:{sc}{last}")
            ws.freeze_panes = f"{col_letter(keys.index('status') + 2)}2"
            ws.auto_filter.ref = f"A1:{end}{max(len(recs) + 1, 2)}"
            counts[tab] = f"=COUNTA('{title}'!A:A)-1"

    # ---- はじめに ----
    intro.sheet_properties.tabColor = INK
    intro.sheet_view.showGridLines = False
    for c, w in zip("ABCD", (3, 20, 70, 10)):
        intro.column_dimensions[c].width = w
    intro["B2"] = INTRO_TITLE
    intro["B2"].font = Font(name=FONT, size=18, bold=True, color=INK)
    intro.row_dimensions[2].height = 30
    intro["B3"] = INTRO_LEAD
    intro["B3"].font = Font(name=FONT, size=10, color=SUB)
    intro.merge_cells("B3:D3")
    intro["B3"].alignment = st["wrap"]
    intro.row_dimensions[3].height = 30
    r = 5
    intro.cell(row=r, column=2, value="状態の色").font = st["bold"]
    r += 1
    meaning = {"提案": "機械の案。読んで、採用・保留・却下を決める", "採用": "使う（空欄も採用の扱い）",
               "保留": "あとで決める（使わない）", "却下": "使わない（行は消さずに残すと、同じ案が出にくい）",
               "反映済み": "変更の提案を書き込んだ"}
    for s_, (bg, fg) in STATUS_STYLE.items():
        a = intro.cell(row=r, column=2, value=s_)
        a.fill, a.font, a.alignment = st["fill"](bg), Font(name=FONT, bold=True, color=fg), st["center"]
        intro.cell(row=r, column=3, value=meaning[s_]).font = st["body"]
        r += 1
    r += 1
    intro.cell(row=r, column=2, value="タブの案内").font = st["bold"]
    r += 1
    for c, h in zip((2, 3, 4), ("タブ", "書くこと", "件数")):
        x = intro.cell(row=r, column=c, value=h)
        x.font, x.fill, x.border = st["bold"], st["fill"]("EFEDE8"), st["box"]
    r += 1
    for tab in TAB_ORDER:
        head_c = GROUPS[TAB_GROUP[tab]][0]
        a = intro.cell(row=r, column=2, value=TABS[tab][0])
        a.fill, a.font, a.border = st["fill"](head_c), st["bold"], st["box"]
        b = intro.cell(row=r, column=3, value=TAB_HELP[tab] + ("（縦型）" if tab in VERTICAL else ""))
        b.font, b.alignment, b.border = st["body"], st["wrap"], st["box"]
        n = intro.cell(row=r, column=4, value=counts[tab])
        n.font, n.alignment, n.border = st["body"], st["center"], st["box"]
        r += 1
    r += 1
    intro.cell(row=r, column=2, value="書き方のきまり").font = st["bold"]
    r += 1
    for line in INTRO_RULES:
        x = intro.cell(row=r, column=2, value="・")
        x.font, x.alignment = st["body"], Alignment(horizontal="right", vertical="top")
        y = intro.cell(row=r, column=3, value=line)
        y.font, y.alignment = st["body"], st["wrap"]
        intro.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        intro.row_dimensions[r].height = 30
        r += 1
    wb.calculation.fullCalcOnLoad = True   # 開いたときに件数の数式を計算し直す
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------------------
# Google スプレッドシート
# ---------------------------------------------------------------------------
class GSheetBook:
    SCOPE = "https://www.googleapis.com/auth/spreadsheets"

    def __init__(self, sheet_id: str, sa_file: Path | None) -> None:
        self.sheet_id = sheet_id
        self.sa = json.loads(Path(sa_file).expanduser().read_text(encoding="utf-8")) if sa_file else None
        self._tok: tuple[str, float] | None = None

    @property
    def read_only(self) -> bool:
        return self.sa is None

    def describe(self) -> str:
        return f"Google スプレッドシート（{self.sheet_id[:8]}…{'・閲覧のみ' if self.read_only else ''}）"

    def token(self) -> str:
        if self._tok and self._tok[1] > time.time() + 60:
            return self._tok[0]
        import base64
        from cryptography.hazmat.primitives import hashes, serialization
        from cryptography.hazmat.primitives.asymmetric import padding

        def b64(b: bytes) -> str:
            return base64.urlsafe_b64encode(b).rstrip(b"=").decode()
        now = int(time.time())
        head = b64(json.dumps({"alg": "RS256", "typ": "JWT"}).encode())
        claim = b64(json.dumps({"iss": self.sa["client_email"], "scope": self.SCOPE,
                                "aud": self.sa.get("token_uri", "https://oauth2.googleapis.com/token"),
                                "iat": now, "exp": now + 3600}).encode())
        key = serialization.load_pem_private_key(self.sa["private_key"].encode(), password=None)
        sig = b64(key.sign(f"{head}.{claim}".encode(), padding.PKCS1v15(), hashes.SHA256()))
        body = urllib.parse.urlencode({"grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer",
                                       "assertion": f"{head}.{claim}.{sig}"}).encode()
        req = urllib.request.Request(self.sa.get("token_uri", "https://oauth2.googleapis.com/token"), data=body)
        with urllib.request.urlopen(req, timeout=60, context=ssl_context()) as r:
            d = json.loads(r.read())
        self._tok = (d["access_token"], time.time() + int(d.get("expires_in", 3600)))
        return self._tok[0]

    def _api(self, method: str, path: str, body=None):
        url = f"https://sheets.googleapis.com/v4/spreadsheets/{self.sheet_id}{path}"
        req = urllib.request.Request(url, method=method, data=json.dumps(body).encode() if body is not None else None,
                                     headers={"Authorization": f"Bearer {self.token()}", "Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=120, context=ssl_context()) as r:
            return json.loads(r.read() or b"{}")

    def _grid(self, title: str) -> list[list]:
        d = self._api("GET", "/values/" + urllib.parse.quote(f"'{title}'") + "?valueRenderOption=FORMATTED_VALUE")
        return d.get("values", [])

    def export_xlsx(self) -> bytes:
        """リンク共有の表を Excel 形式で受け取る（鍵なしで読むとき）"""
        url = f"https://docs.google.com/spreadsheets/d/{self.sheet_id}/export?format=xlsx"
        try:
            with urllib.request.urlopen(url, timeout=120, context=ssl_context()) as r:
                blob = r.read()
        except Exception as e:
            raise RuntimeError(f"Google スプレッドシートを読めませんでした（{e}）。共有を『リンクを知っている全員』"
                               "（閲覧者以上）にするか、サービスアカウントの鍵を設定してください（docs/13）") from e
        if not blob.startswith(b"PK"):
            raise RuntimeError("Google スプレッドシートが共有されていません（ログイン画面が返りました）。"
                               "共有を『リンクを知っている全員』にするか、サービスアカウントを設定してください（docs/13）")
        return blob

    def read(self) -> dict[str, list[dict]]:
        if self.read_only:
            from openpyxl import load_workbook
            return read_workbook(load_workbook(io.BytesIO(self.export_xlsx()), data_only=True))
        ranges = "&".join("ranges=" + urllib.parse.quote(f"'{TABS[t][0]}'") for t in TAB_ORDER)
        d = self._api("GET", f"/values:batchGet?{ranges}&valueRenderOption=FORMATTED_VALUE")
        return {tab: parse_grid(tab, vr.get("values", []))[0] for tab, vr in zip(TAB_ORDER, d.get("valueRanges", []))}

    def _check_writable(self) -> None:
        if self.read_only:
            raise ReadOnlyBook("Google スプレッドシートを閲覧のみで読んでいるので、書き込めません")

    def append(self, tab: str, rows: list[dict]) -> None:
        self._check_writable()
        title = TABS[tab][0]
        items, lay = parse_grid(tab, self._grid(title))
        if lay["orient"] == "row":
            m = header_to_key(tab)
            values = [[to_cell(d.get(m[h][0]), m[h][1]) if h in m else "" for h in lay["heads"]] for d in rows]
            rng = urllib.parse.quote(f"'{title}'!A{lay['header_row']}")
            self._api("POST", f"/values/{rng}:append?valueInputOption=RAW&insertDataOption=INSERT_ROWS", {"values": values})
            return
        col = max([i["_pos"] for i in items] + [1])
        data = []
        for d in rows:
            col += 1
            for k, r in lay["rows"].items():
                data.append({"range": f"'{title}'!{col_letter(col)}{r}", "values": [[to_cell(d.get(k), kind_of(tab, k))]]})
        self._batch(data)

    def update(self, tab: str, pos: int, key: str, value) -> None:
        self.update_many([(tab, pos, key, value)])

    def update_many(self, edits: list[tuple[str, int, str, object]]) -> None:
        self._check_writable()
        lays: dict[str, dict] = {}
        data = []
        for tab, pos, key, value in edits:
            title = TABS[tab][0]
            if tab not in lays:
                lays[tab] = parse_grid(tab, self._grid(title))[1]
            lay = lays[tab]
            if lay["orient"] == "row":
                cell = f"{col_letter(lay['heads'].index(key_to_header(tab)[key]) + 1)}{pos}"
            else:
                cell = f"{col_letter(pos)}{lay['rows'][key]}"
            data.append({"range": f"'{title}'!{cell}", "values": [[to_cell(value, kind_of(tab, key))]]})
        self._batch(data)

    def _batch(self, data: list[dict]) -> None:
        for i in range(0, len(data), 500):
            self._api("POST", "/values:batchUpdate", {"valueInputOption": "RAW", "data": data[i:i + 500]})


def open_book():
    """.env に Google スプレッドシートの ID があればそれ（鍵があれば読み書き、無ければ閲覧のみ）、無ければ Excel ファイル"""
    sid, sa = os.environ.get("ARTIST_BOOK_GSHEET_ID"), os.environ.get("GOOGLE_SERVICE_ACCOUNT_FILE")
    if sid:
        return GSheetBook(sheet_id_of(sid), Path(sa) if sa else None)
    return XlsxBook(Path(os.environ.get("ARTIST_BOOK_XLSX") or ROOT / "artists.xlsx"))


def sheet_id_of(v: str) -> str:
    """URL を貼っても ID だけを取り出す"""
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", v)
    return m.group(1) if m else v.strip()


PENDING = OUT / "artist_book" / "pending.xlsx"


def append_rows(book, tab: str, rows: list[dict]) -> str:
    """管理表に行を足す。Google の表に書けないとき（閲覧のみ）は out/artist_book/pending.xlsx に貯める。足した場所を返す"""
    if not rows:
        return ""
    try:
        book.append(tab, rows)
        return book.describe()
    except ReadOnlyBook:
        pend = XlsxBook(PENDING)
        if not PENDING.exists():
            pend.create()
        pend.append(tab, rows)
        where = PENDING.relative_to(ROOT) if PENDING.is_relative_to(ROOT) else PENDING
        print(f"   （Google の表は閲覧のみで読んでいるので、{TABS[tab][0]} の {len(rows)} 件を "
              f"{where} に貯めました。Google の表に貼るか、サービスアカウントを設定すると直接入ります）")
        return str(where)


def adopted(rows: list[dict]) -> list[dict]:
    return [r for r in rows if str(r.get("status", "")).strip() in ADOPTED]
