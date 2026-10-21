"""
アーティストごとの「いまどこで止まっているか」と、仕組み全体の詰まりを判定する部品（dashboard.py が使う）

組ごとの段階（左から順に進む）：
  1. 名前の確認   2. 写真・ジャケット   3. デビュー曲（Suno → 取り込み → 音量 → ジャケット → DistroKid → 配信）
  4. 毎週の制作   5. 写真の刷新（デビューから 6 か月ごと）
各段階の状態：ok（済み）・run（自動で進行中）・wait（だれかの番）・ng（止まっている・期限切れ）・off（まだ先）
「詰まり」＝左から見て最初の ok でも off でもない段階。担当（あなた・Claude・自動）と、何日止まっているかを出す。
"""
from __future__ import annotations

import os
import shutil
from datetime import date, datetime
from pathlib import Path

from _common import MAIN_LABEL, MAIN_LABEL_NAME, OUT, ROOT, all_artists, load_label, monday_of, production_status, read_json

import _songs

STEP_NAMES = ["名前の確認", "写真・ジャケット", "デビュー曲", "毎週の制作", "写真の刷新"]
YOU, CLAUDE, AUTO = "あなた", "Claude", "自動"


def _d(x) -> date | None:
    try:
        return date.fromisoformat(str(x)[:10]) if x else None
    except ValueError:
        return None


def _step(name, state, who="", detail="", action="", since=None) -> dict:
    return {"name": name, "state": state, "who": who, "detail": detail, "action": action, "since": _d(since)}


def _has_photo(slug: str) -> bool:
    if any((k / f"{slug}_photo.jpg").exists() for k in (ROOT / "debut_kit").glob("*/images")):
        return True
    return any((OUT / "visuals" / slug / "photos").glob("*/*_photo_*.jpg"))


def song_step(song: dict, name: str) -> dict:
    """1 曲の公開の進み具合を、段階 1 つ分の判定にする"""
    st = _songs.state(song)
    s, today = st.get("status", ""), date.today()
    title = f"「{song['title']}」" if song.get("title") else ""
    key = os.environ.get("OPENAI_API_KEY")
    if s == "":
        if song["kind"] == "weekly" and not song.get("style"):
            return _step(name, "run", AUTO, f"{title}ブリーフを作っています", "run_week.py（火曜の朝）", song["week"])
        return _step(name, "wait", YOU, f"{title}Suno で Create → いいねした 1 本を SSD の「楽曲情報」へダウンロード",
                     "python3 scripts/publish.py suno --open", song["week"])
    if s == "imported":
        return _step(name, "run", AUTO, f"{title}取り込み済み。音量を整えています", "", st.get("imported_at"))
    if s == "waiting_ffmpeg":
        return _step(name, "ng", YOU, f"{title}音量を整える道具（ffmpeg）がありません", "brew install ffmpeg", st.get("imported_at"))
    if s == "waiting_cover":
        if song["kind"] == "debut" and not key:
            return _step(name, "ng", YOU, f"{title}ジャケットを作る OpenAI の鍵がありません",
                         "python3 scripts/setup_keys.py --only openai", st.get("imported_at"))
        return _step(name, "run", AUTO, f"{title}ジャケットを作っています", "", st.get("imported_at"))
    if s == "ready":
        since = _d(st.get("ready_at"))
        late = since and (today - since).days >= 3
        return _step(name, "ng" if late else "wait", YOU,
                     f"{title}DistroKid に登録（ダッシュボードの「DistroKid に入力する」で下書き → 最後のボタンはあなた）",
                     f"python3 scripts/publish.py draft {song['id']}", since)
    if s == "registered":
        rel = _d(st.get("release_date"))
        if not (os.environ.get("SPOTIFY_CLIENT_ID") and os.environ.get("SPOTIFY_CLIENT_SECRET")):
            return _step(name, "ng", YOU, f"{title}登録済み。Spotify の鍵がないので配信を確かめられません",
                         "python3 scripts/setup_keys.py --only spotify", st.get("registered_at"))
        if rel and (today - rel).days > 7:
            return _step(name, "ng", YOU, f"{title}配信日（{rel}）から 1 週間たっても Spotify に出ていません。DistroKid の画面で状態を確認",
                         "DistroKid → My Music", st.get("registered_at"))
        return _step(name, "run", AUTO, f"{title}登録済み。配信待ち（配信日 {rel or '—'}）", "毎日 10:15 に確認", st.get("registered_at"))
    if s == "live":
        return _step(name, "ok", "", f"{title}配信中（ISRC {st.get('isrc') or '—'}）", st.get("url", ""), st.get("released"))
    return _step(name, "wait", "", f"{title}{s}")


def artist_steps(a: dict) -> list[dict]:
    slug, today = a["slug"], date.today()
    steps = []
    # 1. 名前の確認
    if a.get("name_status") in ("checked", "registered"):
        steps.append(_step(STEP_NAMES[0], "ok", detail="同じ名前なし"))
    else:
        steps.append(_step(STEP_NAMES[0], "ng", f"{YOU}＋{CLAUDE}", "名前の重複確認が済んでいません", "python3 scripts/check_names.py --apply"))
    # 2. 写真・ジャケット（デビュー曲のジャケット）
    debut = [s for s in _songs.debut_songs() if s["slug"] == slug]
    photo = _has_photo(slug)
    cover = all(_songs.cover_source(s) for s in debut) if debut else True
    if photo and cover:
        steps.append(_step(STEP_NAMES[1], "ok", detail="写真とジャケットあり"))
    else:
        miss = "・".join(x for x, ok in (("写真", photo), ("ジャケット", cover)) if not ok)
        if os.environ.get("OPENAI_API_KEY"):
            steps.append(_step(STEP_NAMES[1], "wait", YOU, f"{miss}がまだありません（Mac で 1 行）", f"python3 scripts/kit_images.py --only {slug}"))
        else:
            steps.append(_step(STEP_NAMES[1], "ng", YOU, f"{miss}を作る OpenAI の鍵がありません", "python3 scripts/setup_keys.py --only openai"))
    # 3. デビュー曲
    if debut:
        steps.append(song_step(debut[0], STEP_NAMES[2]))
    else:
        steps.append(_step(STEP_NAMES[2], "ok", detail="（デビュー曲の作業セットなし）"))
    # 4. 毎週の制作
    ok, why = production_status(a, monday_of(today))
    weekly = sorted([s for s in _songs.weekly_songs() if s["slug"] == slug], key=lambda s: s["week"], reverse=True)
    if not ok and not weekly:
        steps.append(_step(STEP_NAMES[3], "off", detail=why))
    else:
        open_songs = [s for s in weekly if _songs.state(s).get("status") != "live"]
        if open_songs:
            steps.append(song_step(open_songs[-1], STEP_NAMES[3]))   # いちばん古い未完了の曲
        else:
            steps.append(_step(STEP_NAMES[3], "ok" if weekly else "run", AUTO if not weekly else "",
                               f"配信中 {len(weekly)} 曲" if weekly else f"今週から制作（{why}）"))
    # 5. 写真の刷新
    try:
        from refresh_photos import next_due
        due = next_due(slug)
    except Exception:  # noqa: BLE001
        due = None
    if due is None:
        steps.append(_step(STEP_NAMES[4], "off", detail="デビューの配信後に決まる"))
    elif due <= today:
        who, state = (AUTO, "run") if os.environ.get("OPENAI_API_KEY") else (YOU, "ng")
        steps.append(_step(STEP_NAMES[4], state, who, f"刷新の時期です（{due}）。毎月 1 日に自動で作ります",
                           f"python3 scripts/refresh_photos.py --only {slug} --force", due))
    else:
        steps.append(_step(STEP_NAMES[4], "ok", detail=f"次回 {due}"))
    return steps


def blocker(steps: list[dict]) -> dict | None:
    """左から見て最初の、済みでもまだ先でもない段階"""
    return next((s for s in steps if s["state"] not in ("ok", "off")), None)


def artists_overview() -> list[dict]:
    out = []
    for a in all_artists():
        if not a.get("slug"):
            continue
        steps = artist_steps(a)
        b = blocker(steps)
        days = (date.today() - b["since"]).days if b and b["since"] else None
        lab = a.get("label_slug")
        label = MAIN_LABEL_NAME if lab in (None, "", "drive", MAIN_LABEL) else (load_label(lab).get("name") or lab)
        out.append({"slug": a["slug"], "name": a.get("name", a["slug"]), "label": label,
                    "steps": steps, "blocker": b, "days": days})
    order = {"ng": 0, "wait": 1, "run": 2}
    out.sort(key=lambda r: (order.get(r["blocker"]["state"], 3) if r["blocker"] else 4, -(r["days"] or 0)))
    return out


def system_blockers() -> list[tuple[str, str, str]]:
    """仕組み全体の詰まり：(重要度 ng/normal, 内容, 直し方)"""
    out = []
    songs = _songs.all_songs()
    states = [_songs.state(s).get("status") for s in songs]
    if not shutil.which("ffmpeg"):
        out.append(("ng", "音量を整える道具（ffmpeg）が入っていません", "brew install ffmpeg"))
    if not os.environ.get("OPENAI_API_KEY"):
        out.append(("ng", "画像（写真・ジャケット）を作る OpenAI の鍵がありません", "python3 scripts/setup_keys.py --only openai"))
    if not (os.environ.get("SPOTIFY_CLIENT_ID") and os.environ.get("SPOTIFY_CLIENT_SECRET")):
        sev = "ng" if "registered" in states else "normal"
        out.append((sev, "Spotify の鍵がないので、配信の確認と ISRC の記録が自動になりません", "python3 scripts/setup_keys.py --only spotify"))
    from _guide import _has_playwright
    if not _has_playwright():
        out.append(("normal", "Suno・DistroKid に自動で入力する部品（Playwright）が入っていません（無くても、手でコピーして貼れば進められます）",
                    "python3 -m pip install --user playwright && python3 -m playwright install chromium"))
    if not os.environ.get("OWNER_LEGAL_NAME"):
        out.append(("normal", "DistroKid の作詞作曲者の欄に入れる本名が未設定です", "python3 scripts/setup_keys.py --only owner"))
    try:
        import _archive as archive
        base, how = archive.root()
        if base is None:
            n = len(archive.pending())
            out.append(("normal", f"SSD がつながっていません（SSD にまだ写していないもの {n} 件。管理表の書き込みも保留中）", "SSD をつなぐ → python3 scripts/archive_sync.py"))
    except Exception:  # noqa: BLE001
        pass
    agents = Path.home() / "Library" / "LaunchAgents"
    if agents.is_dir() and not (agents / "jp.etherarchi.label.intake.plist").exists():
        out.append(("ng", "定期実行（曲の取り込み・配信の確認）が Mac に登録されていません", "python3 scripts/schedule.py install"))
    logs = OUT / "logs"
    for job, ja in (("intake", "曲の取り込み"), ("release", "配信の確認"), ("backup", "保管"), ("auto", "毎週の制作"), ("brief", "ブリーフ")):
        f = logs / f"{job}.log"
        if not f.exists():
            continue
        last = f.read_text(encoding="utf-8", errors="replace").split("=====")[-1]
        if "フルディスクアクセス" in last:
            out.append(("ng", "自動の処理が「ダウンロード」フォルダを読めません", "システム設定 → プライバシーとセキュリティ → フルディスクアクセス に python3 を追加"))
        elif any(c != "0" for c in __import__("re").findall(r"→ 終了コード (\d+)", last)):
            out.append(("ng", f"定期実行「{ja}」の最後の実行でエラーがありました", f"out/logs/{job}.log を確認"))
    return out


def tracker_rows() -> tuple[list[dict], str]:
    """管理表の「詰まっているところ」タブのうち、完了していないもの。(行, どこから読んだか)"""
    try:
        from _artist_book import ssd_book_path
        p = ssd_book_path(seed=False)
        src = "SSD の管理表"
    except Exception:  # noqa: BLE001
        p = None
    if p is None or not Path(p).exists():
        p, src = ROOT / "artists.xlsx", "手元の見本（SSD が無いので古い可能性あり）"
    try:
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        if "詰まっているところ" not in wb.sheetnames:
            return [], src
        rows = list(wb["詰まっているところ"].iter_rows(values_only=True))
    except Exception:  # noqa: BLE001
        return [], src
    head = [str(h or "") for h in rows[0]] if rows else []
    out = []
    for r in rows[1:]:
        d = {h: ("" if v is None else str(v)) for h, v in zip(head, r)}
        if any(d.values()) and d.get("状態") not in ("完了", "解決"):
            out.append(d)
    return out, src
